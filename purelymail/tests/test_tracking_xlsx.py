"""Tests for safe tracking-workbook updates. Uses a synthetic workbook shaped like the real one."""

from __future__ import annotations

import datetime as dt
import hashlib
import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: E402

import tracking_xlsx as T  # noqa: E402


def make_workbook(path: Path, rows=(("service@bgyhub.com", "已使用"), ("123@bgyhub.com", "已使用"))) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = T.SHEET
    ws.append(T.BASE_HEADERS)
    for r in rows:
        ws.append(list(r) + [None] * (6 - len(r)))
    for c in ws[1]:
        c.font = Font(name="Carlito", bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor="FF1F4E78")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Carlito")
    for letter, width in zip("ABCDEF", (28, 12, 24, 14, 20, 28)):
        ws.column_dimensions[letter].width = width
    dv = DataValidation(type="list", formula1='"已使用,未使用,停用"', allow_blank=False)
    dv.add("B2:B500")
    ws.add_data_validation(dv)
    table = Table(displayName=T.TABLE, ref=f"A1:F{len(rows) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    wb.save(path)


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def values(path: Path) -> list[tuple]:
    return list(openpyxl.load_workbook(path).active.iter_rows(values_only=True))


class TrackingTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.wb = self.dir / "BGYHub邮箱使用记录.xlsx"
        make_workbook(self.wb)
        self.backups = self.dir / "backups"
        self.lock = self.dir / "tracking.lock"

    def tearDown(self):
        self.tmp.cleanup()

    def apply(self, ops):
        return T.apply(self.wb, ops, self.backups, self.lock)

    def test_create_appends_row_and_adds_password_column(self):
        before = values(self.wb)
        res = self.apply([T.RowOp("create", "004@bgyhub.com", "Pw4" + "x" * 21, dt.date(2026, 10, 1))])
        self.assertEqual(res.added, ["004@bgyhub.com"])
        self.assertTrue(res.password_column_added)
        rows = values(self.wb)
        self.assertEqual(rows[0], tuple(T.BASE_HEADERS + ["密码"]))
        for old, new in zip(before[1:], rows[1:3]):  # existing rows unchanged
            self.assertEqual(old + (None,), new)
        email, status, platform, date, payment, note, pw = rows[3]
        self.assertEqual((email, status, platform, payment, note, pw),
                         ("004@bgyhub.com", "未使用", None, None, "Purelymail 创建", "Pw4" + "x" * 21))
        self.assertEqual(date.date(), dt.date(2026, 10, 1))
        ws = openpyxl.load_workbook(self.wb).active
        self.assertEqual(ws.tables[T.TABLE].ref, "A1:G4")
        self.assertEqual([c.name for c in ws.tables[T.TABLE].tableColumns], T.BASE_HEADERS + ["密码"])
        self.assertEqual(str(ws.data_validations.dataValidation[0].sqref), "B2:B500")
        self.assertEqual(ws.column_dimensions["A"].width, 28)
        self.assertEqual(ws.column_dimensions["G"].width, T.PASSWORD_COLUMN_WIDTH)
        self.assertEqual(ws["A4"].font.name, "Carlito")
        self.assertEqual(ws["D4"].number_format, T.DATE_FORMAT)
        self.assertEqual(ws["G1"].fill.fgColor.rgb, "FF1F4E78")  # header style copied
        self.assertEqual(oct(self.wb.stat().st_mode & 0o777), "0o600")

    def test_backup_before_every_change(self):
        original = sha(self.wb)
        r1 = self.apply([T.RowOp("create", "004@bgyhub.com", "a" * 24, dt.date(2026, 10, 1))])
        r2 = self.apply([T.RowOp("reset", "004@bgyhub.com", "b" * 24)])
        self.assertEqual(sha(r1.backup), original)
        self.assertNotEqual(r1.backup, r2.backup)
        self.assertEqual(len(list(self.backups.iterdir())), 2)
        self.assertEqual(oct(r1.backup.stat().st_mode & 0o777), "0o600")

    def test_no_change_means_no_write_and_no_backup(self):
        self.apply([T.RowOp("create", "004@bgyhub.com", "a" * 24, dt.date(2026, 10, 1))])
        before = sha(self.wb)
        res = self.apply([T.RowOp("create", "004@bgyhub.com", "a" * 24, dt.date(2026, 10, 1))])
        self.assertEqual(res.unchanged, ["004@bgyhub.com"])
        self.assertIsNone(res.backup)
        self.assertEqual(sha(self.wb), before)

    def test_existing_row_only_blank_cells_filled_never_duplicated(self):
        make_workbook(self.wb, rows=(("service@bgyhub.com", "已使用"),
                                     ("001@bgyhub.com", "已使用", "Shopify", None, "Visa", "my note")))
        res = self.apply([T.RowOp("import", "001@bgyhub.com", "p" * 24, dt.date(2026, 9, 30))])
        self.assertEqual(res.updated, ["001@bgyhub.com"])
        rows = values(self.wb)
        self.assertEqual(len(rows), 3)
        email, status, platform, date, payment, note, pw = rows[2]
        self.assertEqual((status, platform, payment, note, pw), ("已使用", "Shopify", "Visa", "my note", "p" * 24))
        self.assertEqual(date.date(), dt.date(2026, 9, 30))

    def test_different_existing_password_is_not_overwritten_by_create(self):
        self.apply([T.RowOp("create", "004@bgyhub.com", "a" * 24, dt.date(2026, 10, 1))])
        res = self.apply([T.RowOp("import", "004@bgyhub.com", "b" * 24, dt.date(2026, 10, 1))])
        self.assertEqual(len(res.conflicts), 1)
        self.assertEqual(values(self.wb)[3][6], "a" * 24)

    def test_reset_changes_only_password_cell(self):
        self.apply([T.RowOp("create", "004@bgyhub.com", "a" * 24, dt.date(2026, 10, 1))])
        before = values(self.wb)
        self.apply([T.RowOp("reset", "004@bgyhub.com", "c" * 24)])
        after = values(self.wb)
        self.assertEqual(before[3][:6], after[3][:6])
        self.assertEqual(after[3][6], "c" * 24)
        self.assertEqual(before[:3], after[:3])

    def test_duplicate_rows_in_workbook_are_reported_not_guessed(self):
        make_workbook(self.wb, rows=(("001@bgyhub.com", "已使用"), (" 001@BGYHUB.com ", "停用")))
        res = self.apply([T.RowOp("reset", "001@bgyhub.com", "c" * 24)])
        self.assertEqual(res.conflicts[0][0], "001@bgyhub.com")
        self.assertEqual(values(self.wb)[1][6], None)

    def test_refuses_when_open_in_excel(self):
        before = sha(self.wb)
        (self.dir / ("~$" + self.wb.name)).write_text("lock")
        with self.assertRaises(T.ExcelOpenError):
            self.apply([T.RowOp("create", "004@bgyhub.com", "a" * 24, dt.date(2026, 10, 1))])
        self.assertEqual(sha(self.wb), before)
        self.assertFalse(self.backups.exists() and any(self.backups.iterdir()))

    def test_crash_before_replace_leaves_original_intact(self):
        before = sha(self.wb)
        with mock.patch.object(T.os, "replace", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                self.apply([T.RowOp("create", "004@bgyhub.com", "a" * 24, dt.date(2026, 10, 1))])
        self.assertEqual(sha(self.wb), before)
        self.assertEqual([p for p in self.dir.iterdir() if p.name.startswith(".~bgyhub-tmp-")], [])

    def test_concurrent_modification_aborts(self):
        before_fp = T.fingerprint(self.wb)
        real_fp = T.fingerprint
        calls = {"n": 0}

        def fake_fp(path):
            calls["n"] += 1
            return before_fp if calls["n"] == 1 else (0, 0, "changed")

        with mock.patch.object(T, "fingerprint", side_effect=fake_fp):
            with self.assertRaises(T.TrackingError):
                self.apply([T.RowOp("create", "004@bgyhub.com", "a" * 24, dt.date(2026, 10, 1))])
        self.assertEqual(real_fp(self.wb)[2], before_fp[2])

    def test_verification_failure_leaves_original_intact(self):
        before = sha(self.wb)
        with mock.patch.object(T, "_verify", side_effect=T.TrackingError("verification failed")):
            with self.assertRaises(T.TrackingError):
                self.apply([T.RowOp("create", "004@bgyhub.com", "a" * 24, dt.date(2026, 10, 1))])
        self.assertEqual(sha(self.wb), before)

    def test_verify_detects_changed_existing_cell(self):
        wb = openpyxl.load_workbook(self.wb)
        before = T._snapshot(wb)
        wb.active["A2"] = "tampered"
        with self.assertRaises(T.TrackingError):
            T._verify(before, T._snapshot(wb), {}, {}, {}, {})

    def test_rejects_wrong_layout(self):
        wb = openpyxl.load_workbook(self.wb)
        wb.active["C1"] = "renamed"
        wb.save(self.wb)
        with self.assertRaises(T.TrackingError):
            self.apply([T.RowOp("create", "004@bgyhub.com", "a" * 24, dt.date(2026, 10, 1))])
        self.assertFalse(T.describe(self.wb)["valid"])

    def test_rejects_images(self):
        from openpyxl.drawing.image import Image
        if importlib.util.find_spec("PIL") is None:
            self.skipTest("Pillow not installed")
        png = self.dir / "x.png"
        from PIL import Image as PILImage
        PILImage.new("RGB", (2, 2)).save(png)
        wb = openpyxl.load_workbook(self.wb)
        wb.active.add_image(Image(str(png)), "H2")
        wb.save(self.wb)
        with self.assertRaises(T.TrackingError):
            self.apply([T.RowOp("create", "004@bgyhub.com", "a" * 24, dt.date(2026, 10, 1))])

    def test_dropdown_extended_when_table_grows_past_it(self):
        wb = openpyxl.load_workbook(self.wb)
        wb.active.data_validations.dataValidation[0].sqref = "B2:B3"
        wb.save(self.wb)
        self.apply([T.RowOp("create", f"00{i}@bgyhub.com", "a" * 24, dt.date(2026, 10, 1)) for i in (4, 5)])
        ws = openpyxl.load_workbook(self.wb).active
        self.assertEqual(str(ws.data_validations.dataValidation[0].sqref), "B2:B5")

    def test_macos_privacy_block_is_reported_not_raised(self):
        # macOS returns EPERM for Desktop/Documents/Downloads without permission;
        # Python 3.9's Path.exists()/is_file() raise PermissionError for that.
        err = PermissionError(1, "Operation not permitted")
        with mock.patch.object(T.Path, "exists", side_effect=err), \
             mock.patch.object(T.Path, "is_file", side_effect=err):
            info = T.describe(self.wb)
        self.assertFalse(info["valid"])
        self.assertIn("Privacy", info["reason"])

    def test_find_reports_blocked_folders(self):
        blocked = []
        locked = self.dir / "Documents"
        locked.mkdir()
        real_walk = T.os.walk

        def fake_walk(top, onerror=None):
            if Path(top) == locked:
                onerror(PermissionError(1, "Operation not permitted", str(locked)))
                return iter(())
            return real_walk(top, onerror=onerror)

        with mock.patch.object(T.os, "walk", side_effect=fake_walk):
            found = T.find_candidates([locked, self.dir], blocked)
        self.assertEqual(blocked, [str(locked)])
        self.assertIn(self.wb.name, [Path(f["path"]).name for f in found])

    def test_find_skips_backups(self):
        (self.dir / "backups").mkdir()
        make_workbook(self.dir / "backups" / "old.xlsx")
        self.assertEqual([Path(f["path"]).name for f in T.find_candidates([self.dir])], [self.wb.name])

    def test_find_candidates(self):
        (self.dir / "sub").mkdir()
        other = self.dir / "sub" / "other.xlsx"
        openpyxl.Workbook().save(other)
        found = T.find_candidates([self.dir])
        self.assertEqual([Path(f["path"]).name for f in found], [self.wb.name])


if __name__ == "__main__":
    unittest.main()
